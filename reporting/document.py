"""One report, three formats.

Report content is built once as a list of blocks (headings, paragraphs,
bullet lists, tables, figures) and rendered to Markdown, a self-contained
HTML page (figures embedded), and a Word document. The three files can
therefore never disagree.
"""
from __future__ import annotations

import base64
import html
import math
import struct
from dataclasses import dataclass, field
from pathlib import Path
from typing import List

import pandas as pd


def fmt_value(value: object) -> str:
    if value is None:
        return "—"
    try:
        if pd.isna(value):
            return "—"
    except (TypeError, ValueError):
        pass
    if isinstance(value, float):
        if math.isnan(value):
            return "—"
        return f"{value:.3f}"
    return str(value)


@dataclass
class Block:
    kind: str                 # h1 | h2 | h3 | p | note | bullets | table | figure
    text: str = ""
    items: List[str] = field(default_factory=list)
    frame: pd.DataFrame | None = None
    path: Path | None = None  # figure file


class Report:
    def __init__(self, title: str, subtitle: str = ""):
        self.title = title
        self.subtitle = subtitle
        self.blocks: List[Block] = []

    # ── building ────────────────────────────────────────────────────────
    def h1(self, text: str) -> None:
        self.blocks.append(Block("h1", text))

    def h2(self, text: str) -> None:
        self.blocks.append(Block("h2", text))

    def h3(self, text: str) -> None:
        self.blocks.append(Block("h3", text))

    def p(self, text: str) -> None:
        self.blocks.append(Block("p", text))

    def note(self, title: str, text: str) -> None:
        """A shaded box, e.g. "How to read this chart"."""
        self.blocks.append(Block("note", text=text, items=[title]))

    def bullets(self, items: List[str]) -> None:
        if items:
            self.blocks.append(Block("bullets", items=list(items)))

    def table(self, frame: pd.DataFrame | None, caption: str = "") -> None:
        if frame is None or frame.empty:
            return
        if caption:
            number = 1 + sum(block.kind == "table" for block in self.blocks)
            self.p(f"Table {number}. {caption}")
        self.blocks.append(Block("table", frame=frame.reset_index(drop=True)))

    def figure(self, path: Path | None, caption: str) -> None:
        if path is not None and Path(path).is_file():
            number = 1 + sum(block.kind == "figure" for block in self.blocks)
            self.blocks.append(Block("figure", text=f"Figure {number}. {caption}", path=Path(path)))

    # ── Markdown ────────────────────────────────────────────────────────
    def to_markdown(self, out: Path) -> Path:
        lines = [f"# {self.title}", ""]
        if self.subtitle:
            lines += [self.subtitle, ""]
        for b in self.blocks:
            if b.kind == "h1":
                lines += [f"## {b.text}", ""]
            elif b.kind == "h2":
                lines += [f"### {b.text}", ""]
            elif b.kind == "h3":
                lines += [f"#### {b.text}", ""]
            elif b.kind == "p":
                lines += [b.text, ""]
            elif b.kind == "note":
                lines += [f"> **{b.items[0]}** {b.text}", ""]
            elif b.kind == "bullets":
                lines += [f"- {item}" for item in b.items] + [""]
            elif b.kind == "table":
                cols = list(b.frame.columns)
                lines.append("| " + " | ".join(map(str, cols)) + " |")
                lines.append("|" + "|".join(["---"] * len(cols)) + "|")
                for row in b.frame.itertuples(index=False):
                    cells = [fmt_value(v).replace("|", "\\|").replace("\n", " ") for v in row]
                    lines.append("| " + " | ".join(cells) + " |")
                lines.append("")
            elif b.kind == "figure":
                rel = b.path.relative_to(out.parent).as_posix()
                lines += [f"![{b.text}]({rel})", "", f"*{b.text}*", ""]
        out.write_text("\n".join(lines), encoding="utf-8")
        return out

    # ── HTML (self-contained) ───────────────────────────────────────────
    def to_html(self, out: Path) -> Path:
        e = html.escape
        parts = [
            "<!doctype html><html lang='en'><head><meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            f"<title>{e(self.title)}</title><style>",
            "body{font-family:Arial,Helvetica,sans-serif;max-width:1000px;"
            "margin:32px auto;padding:0 24px;color:#1f2328;background:#fff;line-height:1.55}",
            "h1{font-size:29px;line-height:1.2;border-bottom:3px solid #0969da;padding-bottom:12px}"
            "h2{font-size:22px;margin-top:46px;border-bottom:1px solid #d0d7de;padding-bottom:7px}"
            "h3{font-size:17px;margin-top:30px;color:#1f4d78}h4{font-size:15px;color:#1f4d78}",
            ".sub{color:#57606a;font-size:14px}.tbl{overflow-x:auto;margin:12px 0 26px}",
            ".table-caption{font-size:13px;font-weight:600;color:#1f4d78;margin:18px 0 5px}",
            ".note{background:#eef5fc;border-left:4px solid #0969da;padding:12px 16px;margin:16px 0 24px}",
            "table{border-collapse:collapse;width:100%;font-size:13px}th,td{border-bottom:1px solid #d0d7de;"
            "padding:8px 10px;text-align:left;vertical-align:top}th{background:#e8eef5;color:#17365d}"
            "tbody tr:nth-child(even){background:#f8fafc}"
            "td.n{text-align:right;font-variant-numeric:tabular-nums}",
            "figure{margin:18px 0 32px}img{display:block;max-width:100%;max-height:720px;"
            "height:auto;margin:auto;border:1px solid #eaeef2}"
            "figcaption{color:#57606a;font-size:13px;margin-top:8px}"
            "@media print{body{max-width:none;margin:0;padding:0}figure,.note,tr{break-inside:avoid}"
            "h2,h3,h4{break-after:avoid}}",
            "</style></head><body>",
            f"<h1>{e(self.title)}</h1>",
        ]
        if self.subtitle:
            parts.append(f"<p class='sub'>{e(self.subtitle)}</p>")
        for b in self.blocks:
            if b.kind in ("h1", "h2", "h3"):
                level = {"h1": 2, "h2": 3, "h3": 4}[b.kind]
                parts.append(f"<h{level}>{e(b.text)}</h{level}>")
            elif b.kind == "p":
                cls = " class='table-caption'" if b.text.startswith("Table ") else ""
                parts.append(f"<p{cls}>{e(b.text)}</p>")
            elif b.kind == "note":
                parts.append(f"<div class='note'><b>{e(b.items[0])}</b> {e(b.text)}</div>")
            elif b.kind == "bullets":
                parts.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in b.items) + "</ul>")
            elif b.kind == "table":
                head = "".join(f"<th>{e(str(c))}</th>" for c in b.frame.columns)
                rows = []
                for row in b.frame.itertuples(index=False):
                    cells = "".join(
                        f"<td class='n'>{e(fmt_value(v))}</td>" if isinstance(v, (int, float))
                        else f"<td>{e(fmt_value(v))}</td>"
                        for v in row
                    )
                    rows.append(f"<tr>{cells}</tr>")
                parts.append(f"<div class='tbl'><table><tr>{head}</tr>{''.join(rows)}</table></div>")
            elif b.kind == "figure":
                data = base64.b64encode(b.path.read_bytes()).decode()
                parts.append(
                    f"<figure><img alt='{e(b.text)}' src='data:image/png;base64,{data}'>"
                    f"<figcaption>{e(b.text)}</figcaption></figure>"
                )
        parts.append("</body></html>")
        out.write_text("".join(parts), encoding="utf-8")
        return out

    # ── Word ────────────────────────────────────────────────────────────
    def to_docx(self, out: Path) -> Path:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = section.right_margin = Inches(0.7)
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.header_distance = section.footer_distance = Inches(0.3)
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.12
        for key, size, before, after in (("Title", 22, 0, 9), ("Heading 1", 16, 16, 8),
                                         ("Heading 2", 13, 13, 6), ("Heading 3", 11.5, 10, 5)):
            style = doc.styles[key]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(31, 77, 120)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
        caption_style = doc.styles.add_style("Report Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption_style.font.name = "Calibri"
        caption_style.font.size = Pt(9)
        caption_style.font.italic = True
        caption_style.font.color.rgb = RGBColor(87, 96, 106)
        caption_style.paragraph_format.space_after = Pt(12)
        doc.styles["List Bullet"].paragraph_format.space_after = Pt(4)

        header = section.header.paragraphs[0]
        header.text = "HALLUCINATION BENCHMARK  /  RESEARCH RESULTS"
        header.style = doc.styles["Report Caption"]
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.style = doc.styles["Report Caption"]
        footer.add_run("Page ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

        doc.add_heading(self.title, level=0)
        if self.subtitle:
            doc.add_paragraph(self.subtitle)
        for b in self.blocks:
            if b.kind in ("h1", "h2", "h3"):
                doc.add_heading(b.text, level={"h1": 1, "h2": 2, "h3": 3}[b.kind])
            elif b.kind == "p":
                if b.text.startswith("Table "):
                    para = doc.add_paragraph(b.text, style="Report Caption")
                    para.runs[0].bold = True
                    para.paragraph_format.keep_with_next = True
                else:
                    doc.add_paragraph(b.text)
            elif b.kind == "note":
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Inches(0.16)
                para.paragraph_format.right_indent = Inches(0.16)
                para.paragraph_format.space_before = Pt(6)
                para.paragraph_format.space_after = Pt(10)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEF5FC")
                para._p.get_or_add_pPr().append(shade)
                head = para.add_run(b.items[0] + " ")
                head.bold = True
                body = para.add_run(b.text)
            elif b.kind == "bullets":
                for item in b.items:
                    doc.add_paragraph(item, style="List Bullet")
            elif b.kind == "table":
                frame = b.frame
                table = doc.add_table(rows=1, cols=len(frame.columns))
                table.style = "Table Grid"
                table.autofit = False
                weights = [max(9, min(30, max([len(str(col)),
                                               *(len(fmt_value(v)) for v in frame[col].head(20))]) + 2))
                           for col in frame.columns]
                total = sum(weights)
                widths = [9.5 * weight / total for weight in weights]
                for col, width in zip(table.columns, widths):
                    col.width = Inches(width)
                tbl_width = table._tbl.tblPr.find(qn("w:tblW"))
                if tbl_width is not None:
                    tbl_width.set(qn("w:w"), str(int(9.5 * 1440)))
                    tbl_width.set(qn("w:type"), "dxa")
                for cell, col in zip(table.rows[0].cells, frame.columns):
                    cell.text = str(col)
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8EEF5")
                    cell._tc.get_or_add_tcPr().append(shade)
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in frame.itertuples(index=False):
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = fmt_value(value)
                for row_index, row in enumerate(table.rows):
                    for cell, width in zip(row.cells, widths):
                        cell.width = Inches(width)
                        for para in cell.paragraphs:
                            para.paragraph_format.space_after = Pt(2)
                            for run in para.runs:
                                run.font.size = Pt(8.5)
                                if row_index == 0:
                                    run.bold = True
                if doc.paragraphs:
                    doc.paragraphs[-1].paragraph_format.keep_with_next = True
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif b.kind == "figure":
                with b.path.open("rb") as image_file:
                    header = image_file.read(24)
                aspect = 9.0 / 5.5
                if header[:8] == b"\x89PNG\r\n\x1a\n":
                    width_px, height_px = struct.unpack(">II", header[16:24])
                    aspect = width_px / height_px
                width = min(9.3, 5.55 * aspect)
                image_para = doc.add_paragraph()
                image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_para.paragraph_format.keep_with_next = True
                image_para.add_run().add_picture(str(b.path), width=Inches(width))
                doc.add_paragraph(b.text, style="Report Caption")
        doc.save(str(out))
        return out
